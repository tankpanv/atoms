"""Exercise real generated formats and renderers, including changed-file identity."""
import copy
import io
import tempfile
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch

from agent_harness import validate_plan, TaskLedger, source_digest
from artifact_preview import preview_info, page_image, table, verify_artifacts, artifact_fingerprint
from agent import verify_delivery


def plan(path='dist/report.pptx', fmt='pptx', text='Technical analysis'):
    return {'goal':'Deliver actual user report','application_type':'artifact','task_type':'presentation',
            'design':'Present evidence, analysis and recommendations in the requested format',
            'requirements':[{'id':'R1','description':'Real analysis','acceptance':['Open report and read findings'],'verification':'artifact'}],
            'tasks':[{'id':'T1','title':'Produce report','files':[path],'requirement_ids':['R1'],'depends_on':[],'verification':'Open actual report'}],
            'deliverables':[{'path':path,'format':fmt,'title':'Report','requirement_ids':['R1'],'min_units':1,'content_checks':[text]}]}


class ArtifactDeliveryTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);(self.root/'dist').mkdir()

    async def test_file_only_task_completes_without_code_readme_or_web_bootstrap(self):
        import json,os
        from agent import run_agent
        from test_agent_execution import calls
        from agent_session import AgentSession
        value=plan('dist/report.md','md')
        import shlex
        command = "python -c " + shlex.quote('from pathlib import Path; Path("dist").mkdir(exist_ok=True); Path("dist/report.md").write_text("# Technical analysis\\nActual findings and recommendations")')
        sequence=[calls(('run_shell',{'command':command,'requirement_ids':['R1']})),
                  calls(('update_task',{'id':'T1','status':'done','evidence_ids':['V1']})),
                  {'content':'Report generated with actual findings.'}]
        class Gateway:
            def __init__(self,*args):pass
            async def chat(self,*args,**kwargs):return sequence.pop(0)
        with patch('agent.ensure_workspace',return_value=self.root),patch('agent.project_uid',return_value=os.getuid()),patch('agent.ModelGateway',Gateway),patch('model_catalog.catalog',return_value=[{'id':'test','context':128000}]),patch('agent.scaffold_project',side_effect=AssertionError('No web scaffolding')):
            result=await run_agent(uuid.uuid4(),'Produce report','test',lambda *a,**kw:None,plan=json.dumps(value))
        self.assertIn('Report generated',result['summary'])
        self.assertTrue(AgentSession(self.root).state['completed'])
        self.assertTrue(AgentSession(self.root).state['demo']['ready'])
        self.assertFalse((self.root/'README.md').exists())
        self.assertEqual((self.root/'dist/report.md').read_text().splitlines()[0],'# Technical analysis')

    def test_contract_rejects_ambiguous_paths_and_fake_outputs(self):
        valid=validate_plan(plan());self.assertEqual(valid['commands']['dev'],'')
        for change in ('missing','escape','hidden','wrong_format','unknown_requirement','boolean_units'):
            value=plan();item=value['deliverables'][0]
            if change=='missing':value['deliverables']=[]
            elif change=='escape':item['path']='dist/../report.pptx'
            elif change=='hidden':item['path']='dist/.private/report.pptx'
            elif change=='wrong_format':item['format']='docx'
            elif change=='unknown_requirement':item['requirement_ids']=['R2']
            else:item['min_units']=True
            with self.subTest(change=change),self.assertRaises(ValueError):validate_plan(value)

    async def test_ppt_pages_are_actually_rendered_and_content_verified_without_server(self):
        from pptx import Presentation
        from PIL import Image
        presentation=Presentation()
        for title in ('Technical analysis','Conclusion'):
            slide=presentation.slides.add_slide(presentation.slide_layouts[1]);slide.shapes.title.text=title
            slide.placeholders[1].text='Revenue, architecture and recommendations'
        presentation.save(self.root/'dist/report.pptx')
        value=validate_plan(plan());value['deliverables'][0]['min_units']=2
        with patch('runtime.start_runtime',side_effect=AssertionError('File task must not start services')):
            code,report=await verify_delivery(uuid.uuid4(),self.root,value)
        self.assertEqual(code,0,report)
        info=preview_info(self.root,'dist/report.pptx');self.assertEqual(info['pages'],2)
        png=page_image(self.root,'dist/report.pptx',1)
        with Image.open(io.BytesIO(png)) as image:self.assertGreater(image.width,500)
        before=artifact_fingerprint(self.root,value)
        presentation.slides[0].shapes.title.text='Changed';presentation.save(self.root/'dist/report.pptx')
        self.assertNotEqual(before,artifact_fingerprint(self.root,value))
        self.assertEqual(verify_artifacts(self.root,value)[0],1)
        with self.assertRaises(ValueError):page_image(self.root,'dist/report.pptx',3)

    def test_word_pdf_and_table_previews_contain_actual_content(self):
        from docx import Document
        from openpyxl import Workbook
        from PIL import Image
        document=Document();document.add_heading('Technical analysis',0);document.add_paragraph('Actual conclusions');document.save(self.root/'dist/report.docx')
        word=preview_info(self.root,'dist/report.docx');self.assertEqual(word['kind'],'pages')
        self.assertIn('Technical analysis',word['titles'][0])
        self.assertEqual(verify_artifacts(self.root,plan('dist/report.docx','docx'))[0],0)
        book=Workbook();ws=book.active;ws.title='Revenue';ws.append(['Year','Revenue']);ws.append([2025,100]);ws.append([2026,'=B2*2'])
        other=book.create_sheet('Notes');other.append(['Technical analysis'])
        book.save(self.root/'dist/report.xlsx')
        view=table(self.root,'dist/report.xlsx');self.assertEqual(view['sheets'],['Revenue','Notes']);self.assertEqual(view['rows'][2][1]['value'],'=B2*2')
        self.assertTrue(view['rows'][2][1]['formula']);self.assertEqual(table(self.root,'dist/report.xlsx',1)['rows'][0][0]['value'],'Technical analysis')
        self.assertEqual(verify_artifacts(self.root,plan('dist/report.xlsx','xlsx'))[0],0)
        (self.root/'dist/data.csv').write_text('Year,Revenue\n2025,100\n2026,200\n')
        self.assertEqual(table(self.root,'dist/data.csv',offset=1,limit=1)['rows'],[['2025','100']])
        with self.assertRaises(ValueError):table(self.root,'dist/report.xlsx',4)

    def test_missing_corrupt_or_incomplete_files_cannot_pass(self):
        for content in (None,b'not an office file'):
            if content:(self.root/'dist/report.pptx').write_bytes(content)
            self.assertEqual(verify_artifacts(self.root,plan())[0],1)
        (self.root/'dist/report.md').write_text('# Wrong report\nNo findings')
        self.assertEqual(verify_artifacts(self.root,plan('dist/report.md','md'))[0],1)
        value=validate_plan(plan('dist/report.md','md'));ledger=TaskLedger(self.root,value,'Write report')
        self.assertTrue(any('artifact' in issue for issue in ledger.completion_issues({})))
        self.assertFalse(any('README' in issue for issue in ledger.completion_issues({})))
        (self.root/'dist/report.md').write_text('# Technical analysis\nFindings')
        receipt=ledger.record('artifact_check','real document',0,'passed',source_digest({}),['R1']);ledger.update('T1','done','actual content checked',[receipt])
        self.assertEqual(ledger.completion_issues({}),[])

class ArtifactEndpointTests(unittest.TestCase):
    def test_authenticated_artifact_routes_read_contract_and_actual_pages(self):
        from fastapi.testclient import TestClient
        import main,json
        from auth import current_user
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);(root/'dist').mkdir();(root/'.atoms').mkdir()
            doc=Document();doc.add_heading('Technical analysis',0);doc.save(root/'dist/report.docx')
            (root/'.atoms/task-state.json').write_text(json.dumps({'plan':plan('dist/report.docx','docx')}))
            main.app.dependency_overrides[current_user]=lambda:{'id':uuid.uuid4()}
            self.addCleanup(main.app.dependency_overrides.pop,current_user,None)
            client=TestClient(main.app,base_url='http://localhost');identifier=uuid.uuid4()
            with patch('main.owned_project'),patch('main.ensure_workspace',return_value=root),patch('main.connection') as domain_db:
                domain_db.return_value.__enter__.return_value.execute.return_value.fetchone.return_value=None
                entries=client.get(f'/api/projects/{identifier}/artifacts')
                self.assertEqual(entries.status_code,200,entries.text)
                self.assertEqual(entries.json()['artifacts'][0]['title'],'Report')
                response=client.get(f'/api/projects/{identifier}/artifacts/preview',params={'path':'dist/report.docx'})
                self.assertEqual(response.status_code,200,response.text)
                self.assertEqual(response.json()['kind'],'pages')
                image=client.get(f'/api/projects/{identifier}/artifacts/page',params={'path':'dist/report.docx','page':0})
                self.assertEqual(image.status_code,200,image.text[:100]);self.assertTrue(image.content.startswith(b'\x89PNG'))
                for route in ('preview','page','download'):
                    self.assertEqual(client.get(f'/api/projects/{identifier}/artifacts/{route}',params={'path':'../private.docx'}).status_code,404)
                self.assertEqual(client.get(f'/api/projects/{identifier}/artifacts/page',params={'path':'dist/report.docx','page':-1}).status_code,422)
