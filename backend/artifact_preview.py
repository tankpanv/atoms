"""Real, content-addressed Office/PDF rendering and bounded workbook reading."""
import csv
import hashlib
import io
import json
import math
import os
import shutil
import signal
import subprocess
import tempfile
import threading
import zipfile
import uuid
from pathlib import Path

from project_artifacts import artifact_file, artifacts

MAX_BYTES = 100 * 1024 * 1024
OFFICE = {'.pptx', '.ppt', '.docx', '.doc', '.xls', '.xlsx'}
PDF_FORMATS = OFFICE | {'.pdf'}
TABLE_FORMATS = {'.xlsx', '.csv'}
_conversion_slots = threading.BoundedSemaphore(2)
_pdfium_lock = threading.Lock()


def checksum(path):
    result = hashlib.sha256()
    with path.open('rb') as source:
        for chunk in iter(lambda: source.read(65536), b''):
            result.update(chunk)
    return result.hexdigest()


def checked_file(root, name):
    path = artifact_file(root, name)
    if path.stat().st_size > MAX_BYTES:
        raise ValueError('成果预览支持最大100MB文件；可下载原文件')
    if path.suffix.lower() in {'.pptx','.docx','.xlsx'}:
        try:
            with zipfile.ZipFile(path) as archive:
                if sum(i.file_size for i in archive.infolist())>200*1024*1024:
                    raise ValueError('Office文件解压大小超出预览范围')
        except zipfile.BadZipFile as exc:
            raise ValueError('Office文件损坏，无法打开') from exc
    return path


def preview_pdf(root, name):
    from pypdf import PdfReader
    source = checked_file(root, name)
    if source.suffix.lower()=='.pdf':
        target=source
    else:
        if source.suffix.lower() not in OFFICE:
            raise ValueError('此成果不支持页面预览')
        cache = root/'.atoms/previews'/('office-v1-'+checksum(source))
        for directory in (root/'.atoms',root/'.atoms/previews',cache):
            if directory.is_symlink():raise ValueError('预览缓存目录无效')
        target = cache/'preview.pdf'
        if target.is_symlink():raise ValueError('预览缓存文件无效')
        if not target.is_file():
            cache.mkdir(parents=True,exist_ok=True)
            with _conversion_slots, tempfile.TemporaryDirectory(prefix='atoms-office-') as temporary:
                # Conversion uses a copied input and a private macro-disabled profile,
                # never a shared GUI profile or the user's original document.
                work=Path(temporary); profile=work/'profile'; (profile/'user').mkdir(parents=True)
                (profile/'user/registrymodifications.xcu').write_text('''<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Common/Security"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item><item oor:path="/org.openoffice.Office.Common/Save/Document"><prop oor:name="UpdateDocMode" oor:op="fuse"><value>0</value></prop></item></oor:items>''')
                copied=work/('input'+source.suffix.lower()); shutil.copyfile(source,copied)
                executable=shutil.which('libreoffice') or shutil.which('soffice')
                if not executable:
                    raise ValueError('Office预览转换器未安装，请修复服务环境')
                process=subprocess.Popen([executable,'-env:UserInstallation='+profile.as_uri(),
                    '--headless','--nologo','--nodefault','--norestore','--convert-to','pdf',
                    '--outdir',str(work),str(copied)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,start_new_session=True)
                try:
                    output,_=process.communicate(timeout=90)
                except subprocess.TimeoutExpired as exc:
                    os.killpg(process.pid,signal.SIGKILL);process.communicate()
                    raise ValueError('Office转换超时；请检查文件后重试，原文件已保留') from exc
                converted=work/'input.pdf'
                if process.returncode or not converted.is_file():
                    raise ValueError('Office无法转换此文件为可预览页面')
                doc=PdfReader(converted)
                if not len(doc.pages) or doc.is_encrypted:
                    raise ValueError('Office转换结果为空或已加密')
                staging=cache/(uuid.uuid4().hex+'.pdf')
                shutil.copyfile(converted,staging); staging.replace(target)
    doc=PdfReader(target)
    if doc.is_encrypted or not len(doc.pages):
        raise ValueError('文件已加密或没有可预览页面')
    return target


def page_image(root,name,page):
    import pypdfium2 as pdfium
    pdf=preview_pdf(root,name)
    # PDFium is not thread safe: requests may convert in parallel, but render
    # under one bounded lock. No document scripts/form actions are executed.
    with _pdfium_lock, pdfium.PdfDocument(pdf) as doc:
        if not 0 <= page < len(doc):raise ValueError('页码超出范围')
        current=doc[page]
        try:
            width,height=current.get_size()
            zoom=min(1.7,4096/max(1,width),8192/max(1,height), math.sqrt(8_000_000/max(1,width*height)))
            bitmap=current.render(scale=zoom)
            try:
                image=bitmap.to_pil(); output=io.BytesIO();image.save(output,format='PNG');image.close()
                return output.getvalue()
            finally:bitmap.close()
        finally:current.close()


def table(root,name,sheet=0,offset=0,limit=100,column_offset=0):
    path=checked_file(root,name)
    if path.suffix.lower()=='.csv':
        with path.open(encoding='utf-8-sig',errors='replace',newline='') as stream:
            sample=stream.read(8192);stream.seek(0)
            try: dialect=csv.Sniffer().sniff(sample)
            except csv.Error: dialect=csv.excel
            rows=[];total=0;columns=0
            for index,row in enumerate(csv.reader(stream,dialect)):
                columns=max(columns,len(row));total=index+1
                if offset<=index<offset+limit:rows.append(row[column_offset:column_offset+100])
        if sheet!=0:raise ValueError('工作表不存在')
        return {'kind':'table','sheets':[path.stem],'sheet':0,'total_rows':total,'total_columns':columns,
                'offset':offset,'column_offset':column_offset,'rows':rows}
    if path.suffix.lower()!='.xlsx':raise ValueError('文件不支持交互表格预览')
    from openpyxl import load_workbook
    book=load_workbook(path,read_only=True,data_only=False,keep_links=False)
    cached=load_workbook(path,read_only=True,data_only=True,keep_links=False)
    try:
        if not 0<=sheet<len(book.worksheets):raise ValueError('工作表不存在')
        ws=book.worksheets[sheet];rows=[];uncached=False
        if (ws.max_row or 0)>offset and (ws.max_column or 0)>column_offset:
            options={'min_row':offset+1,'max_row':min(ws.max_row,offset+limit),'min_col':column_offset+1,'max_col':min(ws.max_column,column_offset+100)}
            cached_rows=cached.worksheets[sheet].iter_rows(**options)
            for row,values in zip(ws.iter_rows(**options),cached_rows):
                result=[]
                for cell,computed in zip(row,values):
                    formula=cell.data_type=='f';value=computed.value if formula and computed.value is not None else cell.value
                    uncached |= formula and computed.value is None
                    result.append({'value': str(value) if value is not None else '',
                                   'formula':formula,'expression':str(cell.value) if formula else '', 'number_format':cell.number_format})
                rows.append(result)
        return {'kind':'table','sheets':book.sheetnames,'sheet':sheet,'total_rows':ws.max_row or 0,
                'total_columns':ws.max_column or 0,'offset':offset,'column_offset':column_offset,'rows':rows,
                'note':'部分公式无缓存结果，显示原公式；版式预览会用Calc重新计算。' if uncached else ''}
    finally:book.close();cached.close()



def preview_info(root,name,mode='default'):
    path=checked_file(root,name);suffix=path.suffix.lower(); sha=checksum(path)
    if suffix in PDF_FORMATS and (suffix!='.xlsx' or mode=='pages'):
        from pypdf import PdfReader
        doc=PdfReader(preview_pdf(root,name))
        return {'kind':'pages','pages':len(doc.pages),'sha256':sha,'name':path.name,
                'titles':[(page.extract_text() or '').split('\n')[0][:160] for page in doc.pages[:10000]]}
    if suffix in TABLE_FORMATS:
        return {**table(root,name),'sha256':sha,'name':path.name}
    if suffix in {'.md','.txt','.html','.svg'}:
        if path.stat().st_size>4*1024*1024:raise ValueError('文本预览最大4MB，请下载全文')
        return {'kind':'html' if suffix=='.html' else 'image' if suffix=='.svg' else 'text',
                'text':path.read_text(encoding='utf-8-sig',errors='replace'),'sha256':sha,'name':path.name}
    if suffix in {'.png','.jpg','.jpeg','.webp'}:
        from PIL import Image
        with Image.open(path) as img:img.verify()
        return {'kind':'image','sha256':sha,'name':path.name}
    if suffix=='.zip':
        with zipfile.ZipFile(path) as archive:
            return {'kind':'archive','entries':[i.filename for i in archive.infolist()][:2000],'sha256':sha,'name':path.name}
    raise ValueError('此格式可下载，暂不支持直接预览')


def content(root,name):
    path=checked_file(root,name);suffix=path.suffix.lower()
    if suffix in PDF_FORMATS and suffix!='.xlsx':
        from pypdf import PdfReader
        doc=PdfReader(preview_pdf(root,name))
        return len(doc.pages),'\n'.join(page.extract_text() or '' for page in doc.pages)
    if suffix=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=False,keep_links=False)
        try:
            count=0;values=[];length=0
            for ws in book:
                values.append(ws.title)
                for row in ws.values:
                    if any(v is not None for v in row):count+=1
                    text=' '.join(str(v) for v in row if v is not None);length+=len(text)
                    if length>20_000_000:raise ValueError('表格正文超出验收读取范围')
                    values.append(text)
            return count,'\n'.join(values)
        finally:book.close()
    if suffix=='.csv':
        result=table(root,name);return result['total_rows'],path.read_text(encoding='utf-8-sig',errors='replace')
    info=preview_info(root,name)
    return 1,info.get('text','')


def verify_artifacts(root,plan):
    report=[]
    for item in plan.get('deliverables',[]):
        try:
            info=preview_info(root,item['path']);units,text=content(root,item['path'])
            if info['kind']=='pages':page_image(root,item['path'],0)
            if units<item.get('min_units',1):raise ValueError('页数/数据行数未满足交付要求')
            normalized=''.join(text.split()).casefold()
            missing=[v for v in item.get('content_checks',[]) if ''.join(v.split()).casefold() not in normalized]
            if missing:raise ValueError('成果正文缺少要求内容：'+', '.join(missing))
            report.append(f"{item['path']}：真实文件可打开，{units}页/行，内容要求通过，sha256={info['sha256']}")
        except Exception as exc:
            return 1,'\n'.join(report+[f"{item['path']}：成果验收失败：{exc}"])
    return 0,'\n'.join(report)


def artifact_fingerprint(root,plan):
    result={}
    for item in plan.get('deliverables',[]):
        try:result[item['path']]=checksum(artifact_file(root,item['path']))
        except ValueError:result[item['path']]=None
    return result
