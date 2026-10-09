"""Validated, project-scoped media attachments for agent messages."""

from __future__ import annotations

import base64
import binascii
import io
import uuid
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from fastapi import HTTPException
from pydantic import BaseModel, Field
from PIL import Image, UnidentifiedImageError
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from agent import project_root


MEDIA_TYPES = {
    "image/png": ("png", 5 * 1024 * 1024),
    "image/jpeg": ("jpg", 5 * 1024 * 1024),
    "image/webp": ("webp", 5 * 1024 * 1024),
    "image/gif": ("gif", 5 * 1024 * 1024),
}
DOCUMENT_TYPES = {
    ".txt": "text/plain", ".md": "text/markdown", ".json": "application/json",
    ".csv": "text/csv", ".html": "text/html", ".css": "text/css",
    ".js": "text/javascript", ".jsx": "text/javascript",
    ".ts": "text/plain", ".tsx": "text/plain",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
MAX_DOCUMENT_TEXT = 150_000


class AttachmentInput(BaseModel):
    name: str = Field(min_length=1, max_length=180)
    mime: str
    data: str


def extract_document(name: str, extension: str, content: bytes) -> str:
    try:
        if extension == "pdf":
            reader = PdfReader(io.BytesIO(content))
            if len(reader.pages) > 80:
                raise HTTPException(422, f"{name} 超过 80 页限制")
            extracted = "\n\n".join(f"[第 {index + 1} 页]\n{page.extract_text() or ''}"
                                   for index, page in enumerate(reader.pages))
        elif extension == "docx":
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                    raise HTTPException(422, f"{name} 解压后超过大小限制")
            document = Document(io.BytesIO(content))
            paragraphs = [item.text for item in document.paragraphs]
            for table in document.tables:
                paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            extracted = "\n".join(paragraphs)
        else:
            extracted = content.decode("utf-8-sig")
            if "\x00" in extracted:
                raise ValueError("包含二进制内容")
    except (ValueError, UnicodeError, OSError, BadZipFile, KeyError, PdfReadError, PackageNotFoundError) as exc:
        raise HTTPException(422, f"{name} 文档无法解析：{exc}") from exc
    if not extracted.strip():
        raise HTTPException(422, f"{name} 没有可提取的文字；扫描版文档请先进行 OCR")
    if len(extracted) > MAX_DOCUMENT_TEXT:
        raise HTTPException(413, f"{name} 提取后超过 {MAX_DOCUMENT_TEXT} 字符限制")
    return extracted


def validate_attachments(attachments: list[AttachmentInput]) -> list[tuple[AttachmentInput, bytes, str, str, str]]:
    if len(attachments) > 12:
        raise HTTPException(422, "一次最多上传 4 张图片和 8 份文档")
    validated = []
    total = 0
    image_count = 0
    document_count = 0
    for attachment in attachments:
        spec = MEDIA_TYPES.get(attachment.mime)
        suffix = Path(attachment.name).suffix.lower()
        is_image = spec is not None
        if is_image:
            image_count += 1
            extension, limit = spec
            if suffix not in (f".{extension}", ".jpeg" if extension == "jpg" else f".{extension}"):
                raise HTTPException(422, f"{attachment.name} 扩展名与图片类型不一致")
        else:
            if suffix not in DOCUMENT_TYPES:
                raise HTTPException(422, f"不支持的文档类型：{attachment.name}")
            document_count += 1
            extension = suffix[1:]
            limit = (5 if extension in ("pdf", "docx") else 2) * 1024 * 1024
            expected = DOCUMENT_TYPES[suffix]
            if attachment.mime not in (expected, "application/octet-stream", "", "text/plain"):
                raise HTTPException(422, f"{attachment.name} MIME 类型与扩展名不一致")
        if image_count > 4 or document_count > 8:
            raise HTTPException(422, "一次最多上传 4 张图片和 8 份文档")
        if len(attachment.data) > (limit * 4 // 3) + 8:
            raise HTTPException(413, f"{attachment.name} 超过大小限制")
        try:
            content = base64.b64decode(attachment.data, validate=True)
        except (binascii.Error, ValueError) as exc:
            raise HTTPException(422, f"{attachment.name} 不是有效的文件数据") from exc
        if not content or len(content) > limit:
            raise HTTPException(422, f"{attachment.name} 文件为空或超过大小限制")
        extracted = ""
        if is_image:
            signatures = {
                "image/png": content.startswith(b"\x89PNG\r\n\x1a\n"),
                "image/jpeg": content.startswith(b"\xff\xd8\xff"),
                "image/webp": content.startswith(b"RIFF") and content[8:12] == b"WEBP",
                "image/gif": content.startswith((b"GIF87a", b"GIF89a")),
            }
            if not signatures[attachment.mime]:
                raise HTTPException(422, f"{attachment.name} 图片内容与类型不匹配")
            try:
                with Image.open(io.BytesIO(content)) as image:
                    if image.width * image.height > 40_000_000:
                        raise HTTPException(422, f"{attachment.name} 图片像素过大")
                    image.verify()
            except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError) as exc:
                raise HTTPException(422, f"{attachment.name} 图片数据损坏") from exc
        else:
            extracted = extract_document(attachment.name, extension, content)
        total += len(content)
        if total > 40 * 1024 * 1024:
            raise HTTPException(413, "附件总大小不能超过 40 MB")
        validated.append((attachment, content, extension, "image" if is_image else "document", extracted))
    return validated


def store_attachments(conn, project_id: uuid.UUID, message_id: uuid.UUID, owner_id: uuid.UUID,
                      attachments: list[tuple[AttachmentInput, bytes, str, str, str]]) -> None:
    if not attachments:
        return
    directory = project_root(project_id, owner_id) / ".atoms-attachments"
    directory.mkdir(mode=0o755, parents=True, exist_ok=True)
    for attachment, content, extension, kind, extracted in attachments:
        attachment_id = uuid.uuid4()
        filename = f"{attachment_id.hex}.{extension}"
        path = directory / filename
        path.write_bytes(content)
        path.chmod(0o644)
        conn.execute("""
            INSERT INTO message_attachments(id,project_id,message_id,filename,mime_type,kind,size_bytes,relative_path,extracted_text)
            VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s)
        """, (attachment_id, project_id, message_id, Path(attachment.name).name,
              attachment.mime if kind == "image" else DOCUMENT_TYPES[f".{extension}"], kind,
              len(content), f".atoms-attachments/{filename}", extracted))
